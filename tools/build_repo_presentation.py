"""Build a short presentation describing rnavail and its saved example runs."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "rnavail_repo_presentation.pptx"
ASSETS = ROOT / ".presentation_assets"
ASSETS.mkdir(exist_ok=True)

NAVY = RGBColor(18, 32, 48)
TEAL = RGBColor(13, 148, 136)
MINT = RGBColor(204, 251, 241)
ORANGE = RGBColor(245, 158, 11)
RED = RGBColor(220, 38, 38)
BLUE = RGBColor(37, 99, 235)
INK = RGBColor(31, 41, 55)
MID = RGBColor(82, 97, 112)
PALE = RGBColor(245, 248, 250)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=True):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    return s


def title(slide, heading, kicker=None):
    if kicker:
        add_text(slide, kicker.upper(), 0.72, 0.34, 11.8, 0.28, 10, TEAL, True)
    add_text(slide, heading, 0.72, 0.62, 11.85, 0.58, 28, NAVY, True)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(.72), Inches(1.28), Inches(1.1), Inches(.05))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL; line.line.color.rgb = TEAL


def footer(slide, source):
    add_text(slide, source, .72, 7.16, 11.9, .2, 8, MID)


def bullet_block(slide, items, x, y, w, h, size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.font.name = "Aptos"; p.font.size = Pt(size)
        p.font.color.rgb = color; p.space_after = Pt(10); p.text = "•  " + item
    return box


def add_notes(slide, notes):
    try:
        slide.notes_slide.notes_text_frame.text = notes
    except Exception:
        pass


def chart_metrics():
    path = ASSETS / "mcherry_metrics.png"
    labels = ["Toehold\n(1–26)", "Main stem\n(27–32)"]
    scores = [0.961, 0.641]
    punp = [0.00898, 0.031]
    fig, ax = plt.subplots(figsize=(7.2, 3.3), dpi=180)
    x = range(2)
    ax.bar([i - .18 for i in x], scores, .36, label="Heuristic rank", color="#0d9488")
    ax.bar([i + .18 for i in x], punp, .36, label="Joint P(unpaired)", color="#f59e0b")
    for i, v in enumerate(scores): ax.text(i-.18, v+.03, f"{v:.3f}", ha="center", fontsize=9)
    for i, v in enumerate(punp): ax.text(i+.18, v+.03, f"{v:.3g}", ha="center", fontsize=9)
    ax.set_ylim(0, 1.13); ax.set_xticks(list(x), labels); ax.set_ylabel("0–1 scale")
    ax.spines[['top','right']].set_visible(False); ax.grid(axis='y', alpha=.2)
    ax.legend(frameon=False, loc="upper right")
    fig.tight_layout(); fig.savefig(path, transparent=True); plt.close(fig)
    return path


def crop_report_assets():
    """Make slide-shaped crops from screenshots captured from the real report."""
    overview = Image.open(ASSETS / "report_overview.png")
    overview.crop((285, 105, 1470, 915)).save(ASSETS / "report_overview_crop.png")
    candidate = Image.open(ASSETS / "report_candidate.png")
    candidate.crop((315, 0, 1490, 1040)).save(ASSETS / "report_candidate_crop.png")


def build():
    crop_report_assets()
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    s = prs.slides.add_slide(blank); rect(s, 0, 0, 13.333, 7.5, NAVY, NAVY, False)
    add_text(s, "rnavail", .8, 1.15, 7.8, .8, 42, WHITE, True)
    add_text(s, "Is this RNA target actually available to bind?", .82, 2.02, 9.8, .7, 26, RGBColor(182, 245, 233), True)
    add_text(s, "A practical, provenance-aware screen for RNA target-region accessibility", .82, 2.92, 8.2, .8, 19, WHITE)
    rect(s, 9.7, 1.1, 2.5, 4.6, RGBColor(28, 50, 70), RGBColor(52, 211, 153))
    add_text(s, "SEQUENCE", 10.15, 1.55, 1.6, .3, 11, RGBColor(110, 231, 183), True, PP_ALIGN.CENTER)
    add_text(s, "AUGC…", 10.0, 2.03, 1.9, .45, 25, WHITE, True, PP_ALIGN.CENTER, "Aptos Mono")
    add_text(s, "↓", 10.72, 2.62, .5, .5, 26, RGBColor(110, 231, 183), True, PP_ALIGN.CENTER)
    add_text(s, "OPEN?", 10.1, 3.28, 1.7, .4, 19, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, "P + ΔG", 10.1, 4.15, 1.7, .4, 18, RGBColor(251, 191, 36), True, PP_ALIGN.CENTER)
    add_text(s, "Repository presentation • September 2026", .82, 6.75, 5.5, .25, 10, RGBColor(148, 163, 184))

    # 2
    s = prs.slides.add_slide(blank); title(s, "The sequence can be right — and the target still unreachable", "The problem")
    rect(s, .75, 1.68, 5.55, 4.75, PALE, RGBColor(220, 228, 234))
    add_text(s, "RNA folds back on itself", 1.08, 2.03, 4.8, .45, 23, NAVY, True)
    add_text(s, "A complementary site may already be paired inside a stem. A binder cannot use bases that are structurally occupied.", 1.08, 2.68, 4.65, 1.15, 18, INK)
    add_text(s, "PRESENT ≠ AVAILABLE", 1.08, 4.45, 4.7, .55, 27, RED, True)
    add_text(s, "The molecule samples an ensemble of structures, so “open” is a probability—not a single drawing.", 1.08, 5.25, 4.7, .8, 16, MID)
    rect(s, 6.75, 1.68, 5.8, 4.75, MINT, RGBColor(94, 234, 212))
    add_text(s, "The decision question", 7.15, 2.03, 4.8, .4, 16, TEAL, True)
    add_text(s, "How often is the complete target footprint open at the same time?", 7.15, 2.65, 4.75, 1.25, 27, NAVY, True)
    add_text(s, "Relevant to antisense oligos, toehold triggers, CRISPR spacers, primers and miRNA seeds.", 7.15, 4.45, 4.65, 1.0, 17, INK)
    footer(s, "Source: README.md; docs/01-the-question.md")

    # 3
    s = prs.slides.add_slide(blank); title(s, "Designers need a shortlist—not another plausible-looking number", "The need")
    cards = [
        ("Correct event", "Whole-footprint joint opening, with seed opening kept separate."),
        ("Comparable evidence", "Label model family and scope; do not count duplicate interfaces as independent votes."),
        ("Actionable output", "Scan hundreds of windows cheaply, then deeply evaluate a small non-overlapping shortlist."),
        ("Honest uncertainty", "Surface failures, protocol disagreement, conditioning gaps and missing physics."),
    ]
    for i,(h,b) in enumerate(cards):
        x=.78+(i%2)*6.1; y=1.65+(i//2)*2.45
        rect(s,x,y,5.65,2.0,WHITE,RGBColor(214,223,230))
        add_text(s,str(i+1).zfill(2),x+.3,y+.32,.6,.35,13,TEAL,True)
        add_text(s,h,x+1.0,y+.28,4.2,.4,20,NAVY,True)
        add_text(s,b,x+1.0,y+.85,4.25,.8,15,INK)
    footer(s, "Source: README.md; docs/03-architecture.md; docs/05-consensus-and-scoring.md")

    # 4
    s = prs.slides.add_slide(blank); title(s, "rnavail turns one RNA into ranked, traceable candidates", "The solution")
    add_text(s, "INPUT", .82, 1.72, 1.1, .3, 12, TEAL, True)
    add_text(s, "FASTA or literal sequence", .82, 2.12, 2.35, .6, 20, NAVY, True)
    add_text(s, "+ footprint, recognition,\nconditions, optional probing", .82, 2.85, 2.55, 1.0, 15, MID)
    add_text(s, "→", 3.45, 2.55, .5, .5, 26, TEAL, True, PP_ALIGN.CENTER)
    rect(s, 4.05, 1.65, 5.05, 3.85, MINT, RGBColor(94,234,212))
    add_text(s, "RNAVAIL", 4.45, 2.0, 4.2, .35, 12, TEAL, True, PP_ALIGN.CENTER)
    add_text(s, "Local screen\n+ global confirmation\n+ structural diagnostics", 4.55, 2.62, 4.0, 1.65, 23, NAVY, True, PP_ALIGN.CENTER)
    add_text(s, "One request → compatible readings + provenance", 4.55, 4.65, 4.0, .45, 13, MID, False, PP_ALIGN.CENTER)
    add_text(s, "→", 9.25, 2.55, .5, .5, 26, TEAL, True, PP_ALIGN.CENTER)
    add_text(s, "OUTPUT", 10.0, 1.72, 1.1, .3, 12, TEAL, True)
    add_text(s, "Ranked candidates", 10.0, 2.12, 2.3, .5, 20, NAVY, True)
    bullet_block(s,["joint P(unpaired)", "ΔGopen", "seed accessibility", "warnings + lineage"],10.0,2.8,2.5,2.2,15)
    add_text(s, "Text • JSON • TSV • HTML", .82, 6.18, 11.7, .38, 18, BLUE, True, PP_ALIGN.CENTER)
    footer(s, "Source: README.md; docs/02-pipeline.md; docs/06-outputs.md")

    # 5 workflow
    s = prs.slides.add_slide(blank); title(s, "How it works: screen broadly, confirm deeply", "Pipeline")
    steps=[("1","Declare","Sequence, footprint, event, protocol"),("2","Tile","Every candidate window"),("3","Screen","RNAplfold local joint opening"),("4","Shortlist","Seed-first when mechanism supports it; remove overlap"),("5","Deep stage","Exact global PF + orthogonal diagnostics"),("6","Report","Rank, disagreement, notes, provenance")]
    for i,(n,h,b) in enumerate(steps):
        x=.65+i*2.1
        circ=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x+.52), Inches(1.72), Inches(.72), Inches(.72)); circ.fill.solid(); circ.fill.fore_color.rgb=TEAL; circ.line.color.rgb=TEAL
        add_text(s,n,x+.52,1.86,.72,.28,14,WHITE,True,PP_ALIGN.CENTER)
        if i<5:
            ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+1.26), Inches(2.08), Inches(x+2.09), Inches(2.08)); ln.line.color.rgb=RGBColor(153,173,188); ln.line.width=Pt(2)
        add_text(s,h,x,2.72,1.78,.45,17,NAVY,True,PP_ALIGN.CENTER)
        add_text(s,b,x,3.27,1.78,1.3,12,INK,False,PP_ALIGN.CENTER)
    rect(s,.9,5.25,11.55,1.2,PALE,RGBColor(222,229,235))
    add_text(s,"Key rule",1.18,5.57,1.25,.35,14,TEAL,True)
    add_text(s,"The displayed probability and ΔG come from one coherent opening observation; model disagreement stays visible instead of being averaged away.",2.42,5.46,9.45,.62,17,NAVY,True)
    footer(s, "Source: docs/02-pipeline.md; docs/05-consensus-and-scoring.md")

    # 6 scientific background
    s = prs.slides.add_slide(blank); title(s, "Accessibility is an ensemble property", "A little background")
    rect(s,.78,1.62,3.55,4.95,PALE,RGBColor(214,223,230))
    add_text(s,"1 • RNA is not one shape",1.12,1.98,2.9,.4,19,NAVY,True)
    add_text(s,"The sequence samples many secondary structures. Their relative energies determine how often each structure occurs.",1.12,2.64,2.85,1.1,16,INK)
    add_text(s,"Boltzmann ensemble",1.12,4.17,2.85,.35,15,TEAL,True)
    add_text(s,"low-energy structures\ncarry more statistical weight",1.12,4.72,2.85,.85,17,NAVY,True,PP_ALIGN.CENTER)
    rect(s,4.58,1.62,3.55,4.95,MINT,RGBColor(94,234,212))
    add_text(s,"2 • Ask for an event",4.92,1.98,2.9,.4,19,NAVY,True)
    add_text(s,"For interval i…j, sum the ensemble weight of every structure where all bases in that interval are unpaired.",4.92,2.64,2.85,1.2,16,INK)
    add_text(s,"Pᵤₙₚ(i,j)",4.92,4.25,2.85,.65,32,TEAL,True,PP_ALIGN.CENTER)
    add_text(s,"0 = shut   •   1 = open",4.92,5.12,2.85,.35,15,MID,False,PP_ALIGN.CENTER)
    rect(s,8.38,1.62,4.15,4.95,RGBColor(255,247,237),RGBColor(253,186,116))
    add_text(s,"3 • Convert probability to work",8.72,1.98,3.45,.4,19,NAVY,True)
    add_text(s,"ΔGopen = −RT · ln(Pᵤₙₚ)",8.72,2.82,3.45,.55,23,ORANGE,True,PP_ALIGN.CENTER)
    add_text(s,"At 37 °C, RT ≈ 0.616 kcal/mol",8.72,3.65,3.45,.35,14,MID,False,PP_ALIGN.CENTER)
    add_text(s,"LOW ΔGopen = easier to expose",8.72,4.43,3.45,.4,18,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"ΔG per nucleotide supports comparisons across footprint lengths.",8.72,5.16,3.45,.65,15,INK,False,PP_ALIGN.CENTER)
    footer(s, "Source: docs/01-the-question.md; docs/04-tools.md")

    # 7 calculations
    s = prs.slides.add_slide(blank); title(s, "Two calculations answer the same event at different scopes", "Core calculations")
    rect(s,.76,1.6,5.78,4.95,MINT,RGBColor(94,234,212))
    add_text(s,"LOCAL SCREEN • RNAplfold",1.12,1.94,4.95,.35,14,TEAL,True)
    add_text(s,"Slide a 200-nt context along the RNA",1.12,2.48,4.95,.45,22,NAVY,True)
    bullet_block(s,["One partition-function sweep gives all candidate intervals", "Default maximum pair span: 150 nt", "Fast enough to screen a full transcript", "Produces joint opening and seed-opening lookup tables"],1.12,3.12,4.92,2.45,15)
    add_text(s,"Approximate scope",1.12,5.88,4.9,.3,13,ORANGE,True,PP_ALIGN.CENTER)
    rect(s,6.8,1.6,5.75,4.95,WHITE,RGBColor(204,214,223))
    add_text(s,"GLOBAL CONFIRMATION • Vienna exact",7.16,1.94,4.95,.35,14,TEAL,True)
    add_text(s,"Fold the whole molecule twice",7.16,2.48,4.95,.45,22,NAVY,True)
    add_text(s,"G₀",7.32,3.3,.7,.45,23,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"unconstrained ensemble",8.05,3.34,3.3,.35,15,INK)
    add_text(s,"G₁",7.32,4.05,.7,.45,23,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"target forced unpaired",8.05,4.09,3.3,.35,15,INK)
    add_text(s,"ΔGopen = G₁ − G₀",7.16,4.95,4.95,.45,23,ORANGE,True,PP_ALIGN.CENTER)
    add_text(s,"More expensive, preferred coherent headline observation when compatible",7.16,5.65,4.95,.55,14,MID,False,PP_ALIGN.CENTER)
    footer(s, "Source: docs/04-tools.md §4.3; docs/02-pipeline.md")

    # 8 tool map
    s = prs.slides.add_slide(blank); title(s, "The tools play different roles—they are not twelve equal votes", "Toolchain")
    groups=[
        ("JOINT ACCESSIBILITY",TEAL,["RNAplfold — local, scalable screen","RNAplfold CLI — parity check, same algorithm","Vienna exact — global constrained PF"]),
        ("ENSEMBLE / CROSS-CHECK",BLUE,["RNAfold — global ensemble + entropy","RNAstructure PF — independent code/parameters","Ensemble sample — Monte Carlo state check"]),
        ("ALTERNATIVE MODELS",ORANGE,["CONTRAfold — learned posterior","EternaFold — measurement-trained posterior","LinearFold — fast single structure"]),
        ("BLIND-SPOT DIAGNOSTICS",RED,["ProbKnot — pseudoknots","G4 scan — quadruplex propensity","Kinwalker — co-transcriptional traps"]),
    ]
    for i,(h,c,items) in enumerate(groups):
        x=.76+(i%2)*6.05; y=1.58+(i//2)*2.55
        rect(s,x,y,5.68,2.13,WHITE,RGBColor(214,223,230))
        add_text(s,h,x+.34,y+.28,4.98,.3,13,c,True)
        bullet_block(s,items,x+.34,y+.76,4.98,1.2,14)
    add_text(s,"Only compatible probability estimands are pooled. Posteriors and single structures remain visible diagnostics.",.82,6.77,11.7,.35,15,NAVY,True,PP_ALIGN.CENTER)
    footer(s, "Source: docs/04-tools.md; docs/05-consensus-and-scoring.md")

    # Tool explanations 1/3
    s = prs.slides.add_slide(blank); title(s, "The main calculators: fast search first, careful confirmation second", "What each library contributes • 1 of 3")
    tools=[
        ("RNAplfold", "What it does", "Estimates how often every short stretch is fully open inside a moving local neighborhood.", "Advantage", "Very fast: one pass scores the whole transcript.", "How rnavail uses it", "First-pass screen, seed search, profile and landscape."),
        ("RNAplfold CLI", "What it does", "Runs the same RNAplfold method through ViennaRNA’s official command-line program.", "Advantage", "Checks that the Python integration is producing the same result.", "How rnavail uses it", "Quality-control duplicate—not a second vote."),
        ("Vienna exact", "What it does", "Folds the complete RNA with and without the chosen region forced open; the energy difference is ΔGopen.", "Advantage", "Direct whole-molecule answer for the exact target interval.", "How rnavail uses it", "Preferred headline opening value for shortlisted sites."),
        ("RNAfold", "What it does", "Describes the whole equilibrium ensemble and its most likely low-energy structure.", "Advantage", "Shows overall folding, uncertainty and per-base pairing.", "How rnavail uses it", "Structural context and sanity checks—not a separate joint-opening vote."),
    ]
    for i,t in enumerate(tools):
        x=.72+(i%2)*6.18; y=1.48+(i//2)*2.72
        rect(s,x,y,5.78,2.37,WHITE,RGBColor(211,221,229)); add_text(s,t[0],x+.3,y+.22,5.15,.35,19,NAVY,True)
        add_text(s,t[1].upper(),x+.3,y+.72,1.25,.24,9,TEAL,True); add_text(s,t[2],x+1.55,y+.67,3.92,.55,12,INK)
        add_text(s,t[3].upper(),x+.3,y+1.32,1.25,.24,9,ORANGE,True); add_text(s,t[4],x+1.55,y+1.27,3.92,.42,12,INK)
        add_text(s,t[5].upper(),x+.3,y+1.84,1.25,.24,9,BLUE,True); add_text(s,t[6],x+1.55,y+1.79,3.92,.42,12,NAVY,True)
    footer(s, "Plain-language summary of docs/04-tools.md. “Library” here means the external engine; rnavail supplies an adapter around it.")

    # Tool explanations 2/3
    s = prs.slides.add_slide(blank); title(s, "Cross-checks: ask whether another model tells the same story", "What each library contributes • 2 of 3")
    tools=[
        ("RNAstructure partition", "Calculates pairing probabilities with a separate software codebase and parameter tables.", "Best independent thermodynamic cross-check; disagreement means genuine model sensitivity.", "Adds per-base probability evidence and disagreement diagnostics."),
        ("Ensemble sample", "Randomly draws thousands of structures in proportion to how likely they are.", "Shows whether a site is exposed, partly exposed or buried across actual sampled states.", "Checks the global Vienna ensemble; rare open events become honest upper bounds."),
        ("LinearFold", "Quickly predicts one likely fold for the entire transcript using beam search.", "Scales well and can reveal long-distance pairing missed by a local window.", "Fast structural warning only; one drawing is not an opening probability."),
        ("CONTRAfold", "Uses a machine-learned scoring model instead of laboratory-derived thermodynamic energies.", "A different model class; agreement with thermodynamic tools is reassuring.", "Provides a confidence-like posterior for comparison, never converted to ΔG."),
    ]
    for i,(h,w,a,u) in enumerate(tools):
        x=.72+(i%2)*6.18; y=1.48+(i//2)*2.72
        rect(s,x,y,5.78,2.37,PALE,RGBColor(211,221,229)); add_text(s,h,x+.3,y+.22,5.15,.35,19,NAVY,True)
        add_text(s,"DOES",x+.3,y+.75,.62,.23,9,TEAL,True); add_text(s,w,x+1.05,y+.68,4.38,.55,12,INK)
        add_text(s,"EDGE",x+.3,y+1.34,.62,.23,9,ORANGE,True); add_text(s,a,x+1.05,y+1.27,4.38,.48,12,INK)
        add_text(s,"USED",x+.3,y+1.88,.62,.23,9,BLUE,True); add_text(s,u,x+1.05,y+1.81,4.38,.43,12,NAVY,True)
    footer(s, "Source: docs/04-tools.md §4.4")

    # Tool explanations 3/3
    s = prs.slides.add_slide(blank); title(s, "Specialists look for structures ordinary folding tools cannot see", "What each library contributes • 3 of 3")
    tools=[
        ("EternaFold", "A learned model trained on roughly one million chemical-mapping measurements.", "Tests whether data-trained parameters agree with conventional folding.", "Posterior diagnostic; kept separate from physical probabilities."),
        ("ProbKnot", "Builds a structure that may contain crossing base pairs called pseudoknots.", "Catches a major blind spot shared by normal nested-structure algorithms.", "If it locks the target, rnavail warns that accessibility may be too optimistic."),
        ("G4Hunter scan", "Searches the letters themselves for G-rich patterns likely to form G-quadruplexes.", "Always available and detects non-standard structures ordinary base-pair matrices omit.", "Warning signal only; it does not produce a corrected opening probability."),
        ("Kinwalker", "Simulates one possible folding path while the RNA is being transcribed from 5′ to 3′.", "Can expose kinetic traps hidden by equilibrium calculations.", "Short constructs only; refuses long transcripts rather than hanging the run."),
    ]
    for i,(h,w,a,u) in enumerate(tools):
        x=.72+(i%2)*6.18; y=1.48+(i//2)*2.72
        rect(s,x,y,5.78,2.37,RGBColor(255,250,246),RGBColor(232,218,205)); add_text(s,h,x+.3,y+.22,5.15,.35,19,NAVY,True)
        add_text(s,"DOES",x+.3,y+.75,.62,.23,9,TEAL,True); add_text(s,w,x+1.05,y+.68,4.38,.55,12,INK)
        add_text(s,"EDGE",x+.3,y+1.34,.62,.23,9,ORANGE,True); add_text(s,a,x+1.05,y+1.27,4.38,.48,12,INK)
        add_text(s,"USED",x+.3,y+1.88,.62,.23,9,BLUE,True); add_text(s,u,x+1.05,y+1.81,4.38,.43,12,NAVY,True)
    footer(s, "Source: docs/04-tools.md §4.4")

    # 9 consensus and score
    s = prs.slides.add_slide(blank); title(s, "Ranking is transparent: four evidence gates, then four criteria", "From calculations to rank")
    add_text(s,"CONSENSUS GATES",.82,1.56,4.2,.3,13,TEAL,True)
    gates=["Deduplicate shared calculation groups","Keep estimands separate","Match probing/conditioning","Require protocol compatibility"]
    for i,g in enumerate(gates):
        rect(s,.82,2.03+i*.79,.46,.46,TEAL,TEAL); add_text(s,str(i+1),.82,2.13+i*.79,.46,.2,11,WHITE,True,PP_ALIGN.CENTER)
        add_text(s,g,1.48,2.05+i*.79,4.1,.42,16,NAVY,True)
    add_text(s,"Median is taken across independent compatible groups; disagreement is retained as spread.",.82,5.42,4.85,.8,14,MID)
    rect(s,5.95,1.55,6.55,4.95,PALE,RGBColor(214,223,230))
    add_text(s,"DEFAULT SCORE",6.32,1.9,5.78,.3,13,TEAL,True)
    rows=[("Seed P(open)","40%","10⁻⁴ → 0.5 • log ramp"),("ΔGopen / nt","30%","1.0 → 0.05 • linear"),("Model spread","15%","4.0 → 0.5 kcal/mol"),("Length spread","15%","0.15 → 0.01 kcal/mol/nt")]
    for i,(m,w,a) in enumerate(rows):
        y=2.48+i*.72
        add_text(s,m,6.35,y,2.0,.3,15,NAVY,True); add_text(s,w,8.48,y,1.0,.3,15,TEAL,True,PP_ALIGN.CENTER); add_text(s,a,9.45,y,2.65,.3,13,INK)
    add_text(s,"score = exp[ Σ wᵢ ln(dᵢ) / Σ wᵢ ]",6.35,5.46,5.8,.4,20,ORANGE,True,PP_ALIGN.CENTER)
    add_text(s,"Geometric mean: one poor required property can sink the rank. Missing criteria reduce coverage; they are not silently scored as zero.",6.35,5.94,5.8,.5,13,MID,False,PP_ALIGN.CENTER)
    footer(s, "Source: docs/05-consensus-and-scoring.md §5.1–5.2")

    # 10 worked rank
    s = prs.slides.add_slide(blank); title(s, "Worked rank calculation: candidate w1814", "Saved transcript run")
    add_text(s,"RAW OBSERVATION",.82,1.56,3.1,.3,13,TEAL,True)
    rect(s,.82,1.98,3.6,3.7,MINT,RGBColor(94,234,212))
    add_text(s,"20 nt • positions 1814–1833",1.16,2.32,2.9,.35,16,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"Seed P = 0.760",1.16,3.05,2.9,.4,23,TEAL,True,PP_ALIGN.CENTER)
    add_text(s,"ΔGopen = 2.62 kcal/mol",1.16,3.75,2.9,.35,18,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"ΔGopen / nt = 0.131",1.16,4.42,2.9,.35,18,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"Joint P = 0.014",1.16,5.02,2.9,.3,15,MID,False,PP_ALIGN.CENTER)
    add_text(s,"→",4.65,3.46,.6,.5,28,TEAL,True,PP_ALIGN.CENTER)
    rect(s,5.52,1.98,3.08,3.7,WHITE,RGBColor(210,220,228))
    add_text(s,"DESIRABILITY",5.85,2.32,2.42,.3,13,TEAL,True,PP_ALIGN.CENTER)
    add_text(s,"Seed",5.85,3.08,1.0,.3,16,NAVY,True); add_text(s,"1.000",7.05,3.08,1.0,.3,18,TEAL,True,PP_ALIGN.RIGHT)
    add_text(s,"Opening",5.85,3.78,1.1,.3,16,NAVY,True); add_text(s,"0.915",7.05,3.78,1.0,.3,18,TEAL,True,PP_ALIGN.RIGHT)
    add_text(s,"Model robustness",5.85,4.48,1.55,.45,14,MID); add_text(s,"missing",7.32,4.51,.72,.3,13,MID,False,PP_ALIGN.RIGHT)
    add_text(s,"Length robustness",5.85,5.05,1.55,.45,14,MID); add_text(s,"missing",7.32,5.08,.72,.3,13,MID,False,PP_ALIGN.RIGHT)
    add_text(s,"→",8.84,3.46,.6,.5,28,TEAL,True,PP_ALIGN.CENTER)
    rect(s,9.7,1.98,2.82,3.7,NAVY,NAVY)
    add_text(s,"WEIGHTED\nGEOMETRIC MEAN",10.0,2.38,2.22,.7,13,RGBColor(110,231,183),True,PP_ALIGN.CENTER)
    add_text(s,"0.963",10.0,3.42,2.22,.75,37,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"coverage 70%",10.0,4.48,2.22,.3,14,RGBColor(203,213,225),False,PP_ALIGN.CENTER)
    add_text(s,"This is a preference score, not a 96.3% binding probability.",10.0,5.04,2.22,.48,12,RGBColor(251,191,36),True,PP_ALIGN.CENTER)
    add_text(s,"Because robustness sweeps were not requested, the score uses only 70% of intended weight. Run --robustness and --length-robustness for a better-founded rank.",.95,6.12,11.3,.55,15,RED,True,PP_ALIGN.CENTER)
    footer(s, "Saved run: runs/20260913-175402/report.json; scoring anchors: docs/05-consensus-and-scoring.md")

    # What affects rank
    s = prs.slides.add_slide(blank); title(s, "Not every tool changes the score—and that is intentional", "How all the information is used")
    rect(s,.72,1.52,3.78,4.98,MINT,RGBColor(94,234,212))
    add_text(s,"DIRECTLY RANKS",1.08,1.87,3.02,.3,13,TEAL,True,PP_ALIGN.CENTER)
    bullet_block(s,["Can an initial seed become exposed?", "How much work per base opens the full site?", "Does the answer survive changes in folding settings?", "Does it survive nearby footprint lengths?"],1.08,2.43,3.02,2.9,16)
    add_text(s,"→ four desirabilities → one geometric mean",1.08,5.65,3.02,.48,13,NAVY,True,PP_ALIGN.CENTER)
    rect(s,4.78,1.52,3.78,4.98,PALE,RGBColor(211,221,229))
    add_text(s,"SELECTS / CROSS-CHECKS",5.14,1.87,3.02,.3,13,BLUE,True,PP_ALIGN.CENTER)
    bullet_block(s,["Exact global result supplies the coherent headline P and ΔG", "Compatible probability tools form a de-duplicated consensus", "Local-versus-global difference measures scope sensitivity", "Sampling and alternative models test the story"],5.14,2.43,3.02,2.9,16)
    add_text(s,"→ confidence and interpretation",5.14,5.65,3.02,.4,13,NAVY,True,PP_ALIGN.CENTER)
    rect(s,8.84,1.52,3.78,4.98,RGBColor(255,247,237),RGBColor(253,186,116))
    add_text(s,"WARNS, BUT DOES NOT SCORE",9.2,1.87,3.02,.3,13,ORANGE,True,PP_ALIGN.CENTER)
    bullet_block(s,["Possible pseudoknot", "Possible G-quadruplex", "Kinetic folding trap", "Mutation or transcript-context sensitivity", "Tool failure or unsupported condition"],9.2,2.43,3.02,2.9,16)
    add_text(s,"→ reasons to test, reject or investigate",9.2,5.65,3.02,.48,13,NAVY,True,PP_ALIGN.CENTER)
    footer(s, "The score stays small and auditable; important non-comparable evidence remains visible instead of being forced into one number.")

    # joint-vs-marginal metric
    s = prs.slides.add_slide(blank); title(s, "The important distinction: each base open vs. all bases open", "Core metric")
    rect(s,.75,1.72,4.0,3.85,WHITE,RGBColor(211,220,228)); add_text(s,"Marginal view",1.1,2.05,3.3,.4,20,NAVY,True)
    add_text(s,"Mean per-base unpaired",1.1,2.78,3.2,.3,14,MID)
    add_text(s,"0.81",1.1,3.22,3.2,.8,42,TEAL,True)
    add_text(s,"Looks 81% open",1.1,4.23,3.2,.4,18,INK)
    rect(s,4.95,1.72,4.0,3.85,MINT,RGBColor(94,234,212)); add_text(s,"Joint event",5.3,2.05,3.3,.4,20,NAVY,True)
    add_text(s,"P(all 12 nt unpaired together)",5.3,2.78,3.2,.3,14,MID)
    add_text(s,"0.00026",5.3,3.22,3.2,.8,42,RED,True)
    add_text(s,"≈ 1 time in 3,800",5.3,4.23,3.2,.4,18,INK)
    add_text(s,"≈3,100×",9.35,2.75,3.1,.8,36,ORANGE,True,PP_ALIGN.CENTER)
    add_text(s,"different answer",9.35,3.62,3.1,.4,18,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"Different nucleotides breathe at different times. A full-footprint binder needs them open simultaneously; a shorter seed is a separate event.",.85,6.0,11.65,.6,17,INK,False,PP_ALIGN.CENTER)
    footer(s, "Repository test molecule, region 9–20; README.md and docs/01-the-question.md")

    # 7 mCherry setup
    s = prs.slides.add_slide(blank); title(s, "Worked example: a toehold switch with an mCherry off-target", "Saved run")
    add_text(s,"TARGET",.82,1.68,1.2,.3,11,TEAL,True); add_text(s,"Candidate 1 switch",.82,2.03,3.0,.4,23,NAVY,True)
    add_text(s,"95 nt • GC 52.6%",.82,2.52,2.6,.3,14,MID)
    rect(s,.85,3.05,3.2,1.6,MINT,RGBColor(94,234,212)); add_text(s,"TOEHOLD",1.1,3.34,2.7,.3,12,TEAL,True,PP_ALIGN.CENTER); add_text(s,"positions 1–26",1.1,3.86,2.7,.35,20,NAVY,True,PP_ALIGN.CENTER)
    rect(s,.85,4.92,3.2,1.05,PALE,RGBColor(213,222,229)); add_text(s,"MAIN STEM • 27–32",1.1,5.25,2.7,.3,15,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,"+",4.38,3.58,.5,.5,28,TEAL,True,PP_ALIGN.CENTER)
    rect(s,5.15,2.6,3.25,2.4,WHITE,RGBColor(211,220,228)); add_text(s,"TRIGGER",5.5,2.95,2.55,.3,12,TEAL,True,PP_ALIGN.CENTER); add_text(s,"35 nt",5.5,3.5,2.55,.45,25,NAVY,True,PP_ALIGN.CENTER); add_text(s,"seed length 8",5.5,4.23,2.55,.3,14,MID,False,PP_ALIGN.CENTER)
    add_text(s,"+",8.66,3.58,.5,.5,28,TEAL,True,PP_ALIGN.CENTER)
    rect(s,9.4,2.6,3.1,2.4,RGBColor(255,247,237),RGBColor(253,186,116)); add_text(s,"COMPETITOR",9.72,2.95,2.45,.3,12,ORANGE,True,PP_ALIGN.CENTER); add_text(s,"mCherry",9.72,3.5,2.45,.45,25,NAVY,True,PP_ALIGN.CENTER); add_text(s,"off-target sequence",9.72,4.23,2.45,.3,14,MID,False,PP_ALIGN.CENTER)
    add_text(s,"Run protocol: 37 °C • Turner 2004 • 12 tools ran • 3.3 s",.82,6.38,11.7,.4,17,BLUE,True,PP_ALIGN.CENTER)
    footer(s, "Saved run: runs/20260905-133646-candidate1-alltools/meta.json")

    # 8 results
    s = prs.slides.add_slide(blank); title(s, "mCherry run: the toehold ranks strongly, but “open” is nuanced", "Results")
    chart=chart_metrics(); s.shapes.add_picture(str(chart), Inches(.72), Inches(1.52), width=Inches(6.35))
    rect(s,7.35,1.72,5.25,4.75,PALE,RGBColor(216,224,231))
    add_text(s,"TOEHOLD • 26 nt",7.75,2.02,4.45,.35,14,TEAL,True)
    add_text(s,"Rank 0.961",7.75,2.52,4.45,.45,25,NAVY,True)
    bullet_block(s,["Joint P(open) = 0.00898", "ΔGopen = 2.90 kcal/mol", "Best 8-nt seed P(open) = 0.211", "Saved interaction-stage estimate: ΔGtotal = −29.67 kcal/mol"],7.75,3.15,4.35,2.05,15)
    add_text(s,"Interpretation",7.75,5.38,1.5,.3,13,ORANGE,True)
    add_text(s,"The full footprint is rarely pre-open, but seed nucleation is much more plausible. The rank is a heuristic—not a 96% success probability.",7.75,5.72,4.3,.62,14,INK)
    footer(s, "Saved run: runs/20260905-133646-candidate1-alltools/report.txt. Historical interaction metrics are no longer in the current single-molecule scope.")

    # 9 current run
    s = prs.slides.add_slide(blank); title(s, "A current transcript scan shows the intended workflow at scale", "Run evidence")
    add_text(s,"1,833 nt",.82,1.66,2.3,.6,31,TEAL,True); add_text(s,"transcript",.82,2.2,2.3,.35,16,MID)
    add_text(s,"1,814",3.45,1.66,2.3,.6,31,TEAL,True); add_text(s,"20-nt windows screened",3.45,2.2,2.6,.35,16,MID)
    add_text(s,"10",6.35,1.66,2.3,.6,31,TEAL,True); add_text(s,"candidates deep-tested",6.35,2.2,2.6,.35,16,MID)
    add_text(s,"12",9.25,1.66,2.3,.6,31,TEAL,True); add_text(s,"tools completed",9.25,2.2,2.6,.35,16,MID)
    rect(s,.82,3.0,11.7,2.7,WHITE,RGBColor(214,223,230))
    cols=[("w1814","0.963","0.760","0.014","2.62"),("w344","0.929","0.835","0.00156","3.98"),("w150","0.906","0.314","0.00230","3.74")]
    headers=["Candidate","Rank","Seed P","Joint P","ΔGopen"]
    for j,h in enumerate(headers): add_text(s,h,1.1+j*2.15,3.35,1.9,.3,13,MID,True)
    for i,row in enumerate(cols):
        for j,v in enumerate(row): add_text(s,v,1.1+j*2.15,3.95+i*.5,1.9,.3,15,NAVY,j==0)
    add_text(s,"Important warning: local-window and global calculations differed by 1.6–5.5 kcal/mol for several candidates; one site also carried a pseudoknot warning.",.95,6.05,11.3,.62,16,RED,True,PP_ALIGN.CENTER)
    footer(s, "Saved run: runs/20260913-175402/report.txt and meta.json")

    # HTML report overview screenshot
    s = prs.slides.add_slide(blank); title(s, "What the real HTML report looks like: transcript overview", "Actual output")
    rect(s,.63,1.38,12.05,5.65,WHITE,RGBColor(196,207,216))
    s.shapes.add_picture(str(ASSETS / "report_overview_crop.png"), Inches(.75), Inches(1.5), width=Inches(8.3), height=Inches(5.28))
    add_text(s,"HOW TO READ IT",9.35,1.72,2.85,.3,13,TEAL,True)
    rect(s,9.29,2.25,.52,.52,TEAL,TEAL); add_text(s,"1",9.29,2.35,.52,.2,17,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"Profile",9.98,2.27,2.0,.3,17,NAVY,True)
    add_text(s,"Low valleys cost less energy to open. Shaded bands are shortlisted locations.",9.35,2.76,2.75,.78,14,INK)
    rect(s,9.29,3.72,.52,.52,TEAL,TEAL); add_text(s,"2",9.29,3.82,.52,.2,17,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"Landscape",9.98,3.74,2.0,.3,17,NAVY,True)
    add_text(s,"Shows whether a location remains accessible when the binder length changes.",9.35,4.22,2.75,.72,14,INK)
    rect(s,9.29,5.13,.52,.52,TEAL,TEAL); add_text(s,"3",9.29,5.23,.52,.2,17,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"Ranked table",9.98,5.15,2.0,.3,17,NAVY,True)
    add_text(s,"Links every proposed site to its detailed evidence card.",9.35,5.63,2.75,.62,14,INK)
    footer(s, "Screenshot of runs/20260913-175402/report.html — the report is self-contained and can be opened in any browser.")

    # HTML report candidate screenshot
    s = prs.slides.add_slide(blank); title(s, "What the real HTML report looks like: one candidate explained", "Actual output")
    rect(s,.62,1.38,8.88,5.66,WHITE,RGBColor(196,207,216))
    s.shapes.add_picture(str(ASSETS / "report_candidate_crop.png"), Inches(.72), Inches(1.49), width=Inches(6.38), height=Inches(5.42))
    add_text(s,"THE CARD ANSWERS",7.38,1.67,1.72,.3,12,TEAL,True)
    bullet_block(s,["What are the physical values?", "Which tool supplied the headline?", "What warning changes interpretation?", "Exactly how was the rank built?", "What pairing pattern and fold are predicted?"],7.38,2.22,1.75,3.75,13)
    rect(s,9.78,1.38,2.9,5.66,NAVY,NAVY)
    add_text(s,"Example w1814",10.12,1.76,2.2,.35,17,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"Full site open",10.12,2.48,2.2,.3,13,RGBColor(203,213,225),False,PP_ALIGN.CENTER)
    add_text(s,"1.4%",10.12,2.85,2.2,.58,31,RGBColor(110,231,183),True,PP_ALIGN.CENTER)
    add_text(s,"Seed open",10.12,3.58,2.2,.3,13,RGBColor(203,213,225),False,PP_ALIGN.CENTER)
    add_text(s,"76.0%",10.12,3.95,2.2,.58,31,RGBColor(110,231,183),True,PP_ALIGN.CENTER)
    add_text(s,"The report immediately explains the gap: nucleation may be plausible even when the whole footprint is rarely pre-open.",10.12,4.88,2.2,1.3,14,WHITE,False,PP_ALIGN.CENTER)
    footer(s, "Screenshot of runs/20260913-175402/report.html, candidate w1814.")

    # 10 caveats
    s = prs.slides.add_slide(blank); title(s, "Use the ranking as a shortlist, not a binding prediction", "Limitations & caveats")
    left=["Single-molecule secondary structure only", "No partner binding, concentration, specificity or cellular endpoint", "Equilibrium-focused; kinetics mostly absent", "Mg²⁺, pH, crowding, proteins, ribosomes and 3D contacts are not modeled"]
    right=["Most joint-opening estimates share ViennaRNA lineage", "Local vs. global scope can disagree", "Pseudoknots/G-quadruplexes need special diagnostics", "Window length, transcript boundaries, isoform and variants can reorder results", "Heuristic weights are not experimentally calibrated"]
    rect(s,.78,1.63,5.8,4.95,RGBColor(255,247,237),RGBColor(253,186,116)); add_text(s,"BOUNDARY OF THE CLAIM",1.15,1.98,4.95,.3,13,ORANGE,True); bullet_block(s,left,1.15,2.55,4.95,3.45,16)
    rect(s,6.82,1.63,5.72,4.95,PALE,RGBColor(210,220,228)); add_text(s,"MODEL & INPUT SENSITIVITY",7.18,1.98,4.85,.3,13,TEAL,True); bullet_block(s,right,7.18,2.55,4.85,3.45,16)
    footer(s, "Source: docs/07-limitations.md")

    # 11 future + summary
    s = prs.slides.add_slide(blank); title(s, "Next: connect structural accessibility to measured performance", "Future steps + summary")
    future=[("1","Validate","Test accessible, buried and disagreement candidates in matched conditions."),("2","Calibrate","Fit assay-specific rank/endpoint models on held-out transcripts."),("3","Extend","Optional partner-specific interaction command with chemistry checks."),("4","Contextualize","Isoform/haplotype analysis and condition-matched biological evidence."),("5","Research","Kinetic and CGMD modules only with direct experimental validation.")]
    for i,(n,h,b) in enumerate(future):
        y=1.52+i*.92
        rect(s,.78,y,.5,.5,TEAL,TEAL); add_text(s,n,.78,y+.1,.5,.2,12,WHITE,True,PP_ALIGN.CENTER)
        add_text(s,h,1.48,y-.02,1.25,.35,16,NAVY,True); add_text(s,b,2.82,y-.02,5.7,.55,14,INK)
    rect(s,8.82,1.55,3.7,4.75,NAVY,NAVY)
    add_text(s,"TAKEAWAY",9.2,1.95,2.95,.3,12,RGBColor(110,231,183),True,PP_ALIGN.CENTER)
    add_text(s,"Sequence presence\nis not accessibility.",9.2,2.62,2.95,1.2,24,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"rnavail identifies structurally plausible targets, separates seed from full-footprint opening, and preserves the caveats needed to interpret the rank.",9.2,4.15,2.95,1.35,15,RGBColor(226,232,240),False,PP_ALIGN.CENTER)
    footer(s, "Source: docs/09-rna-accessibility-implementation-proposal.md; README.md")

    # metadata
    prs.core_properties.title = "rnavail — repository overview"
    prs.core_properties.subject = "Problem, need, solution, workflow, saved mCherry example, caveats and future steps"
    prs.core_properties.author = "OpenAI Codex"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
